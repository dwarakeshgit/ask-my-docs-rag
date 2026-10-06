from openpyxl import load_workbook

from .models import DocumentPage


def read_excel(file_path: str) -> list[DocumentPage]:
    workbook = load_workbook(file_path, data_only=True)

    pages = []

    for sheet in workbook.worksheets:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            values = [
                str(value).strip()
                for value in row
                if value is not None
            ]

            if values:
                rows.append(" | ".join(values))

        text = "\n".join(rows)

        if text:
            pages.append(
                DocumentPage(
                    text=text,
                    source=file_path,
                    file_type="xlsx",
                    page_number=None,
                    metadata={
                        "sheet": sheet.title
                    },
                )
            )

    workbook.close()

    return pages