from pptx import Presentation

from .models import DocumentPage


def read_pptx(file_path: str) -> list[DocumentPage]:
    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    text_parts.append(text)

        full_text = "\n".join(text_parts)

        if full_text:
            slides.append(
                DocumentPage(
                    text=full_text,
                    source=file_path,
                    file_type="pptx",
                    page_number=slide_number,
                    metadata={
                        "slide_number": slide_number
                    },
                )
            )

    return slides